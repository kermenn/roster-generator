"""
Excel Import and Export Service for Roster Generator.
Accurately replicates the styling, colors, headers, and formulas of Contoh Roster.xlsx.
"""
import io
import calendar
from datetime import date, datetime
from typing import Dict, List, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from engine.models import SHIFTS, INDONESIAN_DAYS, INDONESIAN_MONTHS
from engine.roster_engine import RosterEngine

# Styling constants matching Contoh Roster.xlsx
THIN_BORDER = Border(
    left=Side(style='thin', color='A0A0A0'),
    right=Side(style='thin', color='A0A0A0'),
    top=Side(style='thin', color='A0A0A0'),
    bottom=Side(style='thin', color='A0A0A0')
)

HEADER_FILL = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
DATE_HEADER_FILL = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
SUMMARY_LABEL_FILL = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

# Shift Fills
FILL_D = PatternFill(start_color='92D050', end_color='92D050', fill_type='solid')    # Green
FILL_DN = PatternFill(start_color='A9D08E', end_color='A9D08E', fill_type='solid')   # Soft Green
FILL_N = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')    # Orange
FILL_OFF = PatternFill(start_color='FF7C80', end_color='FF7C80', fill_type='solid')  # Red/Coral
FILL_FB = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')   # Yellow
FILL_CT = PatternFill(start_color='00B0F0', end_color='00B0F0', fill_type='solid')   # Cyan

SHIFT_FILLS = {
    "D": FILL_D,
    "DN": FILL_DN,
    "N": FILL_N,
    "OFF": FILL_OFF,
    "FB": FILL_FB,
    "CT": FILL_CT
}

ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_LEFT = Alignment(horizontal='left', vertical='center')
FONT_BOLD = Font(name='Calibri', size=11, bold=True)
FONT_REGULAR = Font(name='Calibri', size=11)
FONT_SMALL = Font(name='Calibri', size=9)


class ExcelService:
    @staticmethod
    def export_monthly_excel(
        year: int,
        month: int,
        employees: List[Dict[str, Any]],
        overrides: Dict[str, Any],
        min_standby_counterman: int = 2
    ) -> io.BytesIO:
        """
        Creates an Excel workbook matching the layout of Contoh Roster.xlsx.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Roster {INDONESIAN_MONTHS[month]} {year}"
        ws.views.sheetView[0].showGridLines = True

        num_days = calendar.monthrange(year, month)[1]
        last_day_col = 6 + num_days  # G is col 7, so col 6 + num_days

        # 1. Setup Static Info Headers (Columns A to F)
        info_headers = [
            (1, "No"),
            (2, "Name"),
            (3, "SN"),
            (4, "Position"),
            (5, "Location"),
            (6, "Phone No")
        ]

        for col_idx, label in info_headers:
            col_letter = get_column_letter(col_idx)
            ws.merge_cells(f"{col_letter}1:{col_letter}3")
            cell = ws.cell(1, col_idx, label)
            cell.font = FONT_BOLD
            cell.alignment = ALIGN_CENTER
            cell.fill = HEADER_FILL

            # Apply borders to merged area
            for r in range(1, 4):
                ws.cell(r, col_idx).border = THIN_BORDER

        # 2. Setup Date Headers (Row 1: Month/Year, Row 2: Day Name, Row 3: Day Number)
        start_date_col_letter = get_column_letter(7)
        end_date_col_letter = get_column_letter(last_day_col)

        # Merge Row 1 for Month/Year
        ws.merge_cells(f"{start_date_col_letter}1:{end_date_col_letter}1")
        header_date_cell = ws.cell(1, 7, datetime(year, month, 1))
        header_date_cell.number_format = 'mmmm yyyy'
        header_date_cell.font = Font(name='Calibri', size=12, bold=True)
        header_date_cell.alignment = ALIGN_CENTER
        header_date_cell.fill = DATE_HEADER_FILL

        for c in range(7, last_day_col + 1):
            ws.cell(1, c).border = THIN_BORDER

        # Row 2 & 3: Days & Dates
        days_meta = RosterEngine.get_month_days(year, month)
        for idx, d_info in enumerate(days_meta):
            col_idx = 7 + idx
            # Day name (Row 2)
            c2 = ws.cell(2, col_idx, d_info["short_name"])
            c2.font = FONT_BOLD
            c2.alignment = ALIGN_CENTER
            c2.border = THIN_BORDER
            if d_info["is_weekend"]:
                c2.fill = PatternFill(start_color='FFD9D9', end_color='FFD9D9', fill_type='solid')
            else:
                c2.fill = HEADER_FILL

            # Day number (Row 3)
            c3 = ws.cell(3, col_idx, d_info["day"])
            c3.font = FONT_BOLD
            c3.alignment = ALIGN_CENTER
            c3.border = THIN_BORDER
            if d_info["is_weekend"]:
                c3.fill = PatternFill(start_color='FFD9D9', end_color='FFD9D9', fill_type='solid')
            else:
                c3.fill = HEADER_FILL

        # Extra Summary Columns at the right (Total D, DN, N, Jam Lembur, OFF)
        summary_cols = [
            (last_day_col + 1, "Total D"),
            (last_day_col + 2, "Total DN"),
            (last_day_col + 3, "Total N"),
            (last_day_col + 4, "Lembur (Jam)"),
            (last_day_col + 5, "Total OFF")
        ]
        for col_idx, label in summary_cols:
            col_letter = get_column_letter(col_idx)
            ws.merge_cells(f"{col_letter}1:{col_letter}3")
            cell = ws.cell(1, col_idx, label)
            cell.font = FONT_BOLD
            cell.alignment = ALIGN_CENTER
            cell.fill = HEADER_FILL
            for r in range(1, 4):
                ws.cell(r, col_idx).border = THIN_BORDER

        # 3. Populate Employee Rows
        start_emp_row = 4
        counterman_rows = []

        for emp_idx, emp in enumerate(employees):
            current_row = start_emp_row + emp_idx
            is_counterman = "counterman" in str(emp.get("position", "")).lower()
            if is_counterman:
                counterman_rows.append(current_row)

            # Static info
            ws.cell(current_row, 1, emp.get("no", emp_idx + 1)).alignment = ALIGN_CENTER
            ws.cell(current_row, 2, emp.get("name", "")).alignment = ALIGN_LEFT
            ws.cell(current_row, 3, emp.get("sn", "")).alignment = ALIGN_CENTER
            ws.cell(current_row, 4, emp.get("position", "")).alignment = ALIGN_LEFT
            ws.cell(current_row, 5, emp.get("location", "")).alignment = ALIGN_LEFT
            ws.cell(current_row, 6, emp.get("phone", "")).alignment = ALIGN_CENTER

            for c in range(1, 7):
                ws.cell(current_row, c).font = FONT_REGULAR
                ws.cell(current_row, c).border = THIN_BORDER

            # Shifts for each day
            for day_idx in range(num_days):
                day_num = day_idx + 1
                dt = date(year, month, day_num)
                shift_res = RosterEngine.get_effective_shift(emp["id"], dt, emp, overrides)
                code = shift_res["code"]

                col_idx = 7 + day_idx
                cell = ws.cell(current_row, col_idx, code)
                cell.alignment = ALIGN_CENTER
                cell.font = FONT_BOLD
                cell.border = THIN_BORDER

                # Apply shift color fill
                if code in SHIFT_FILLS:
                    cell.fill = SHIFT_FILLS[code]

            # Right summary formulas
            start_col_letter = get_column_letter(7)
            end_col_letter = get_column_letter(last_day_col)
            range_str = f"{start_col_letter}{current_row}:{end_col_letter}{current_row}"

            # Total D
            ws.cell(current_row, last_day_col + 1, f'=COUNTIF({range_str}, "D")').alignment = ALIGN_CENTER
            # Total DN
            ws.cell(current_row, last_day_col + 2, f'=COUNTIF({range_str}, "DN")').alignment = ALIGN_CENTER
            # Total N
            ws.cell(current_row, last_day_col + 3, f'=COUNTIF({range_str}, "N")').alignment = ALIGN_CENTER
            # Lembur (D*3 + N*3)
            tot_d_col = get_column_letter(last_day_col + 1)
            tot_n_col = get_column_letter(last_day_col + 3)
            ws.cell(current_row, last_day_col + 4, f'=({tot_d_col}{current_row}*3)+({tot_n_col}{current_row}*3)').alignment = ALIGN_CENTER
            # Total OFF
            ws.cell(current_row, last_day_col + 5, f'=COUNTIF({range_str}, "OFF")').alignment = ALIGN_CENTER

            for c in range(last_day_col + 1, last_day_col + 6):
                ws.cell(current_row, c).font = FONT_BOLD
                ws.cell(current_row, c).border = THIN_BORDER

        # 4. Summary Rows at Bottom (Standby, Normal, Off, Total)
        last_emp_row = start_emp_row + len(employees) - 1
        summary_labels = [
            ("Standby", "Standby (D)", "D", counterman_rows if counterman_rows else list(range(start_emp_row, last_emp_row + 1))),
            ("Normal", "Normal (DN)", "DN", list(range(start_emp_row, last_emp_row + 1))),
            ("Off", "Off (OFF)", "OFF", list(range(start_emp_row, last_emp_row + 1))),
            ("Total", "Total", "SUM", None)
        ]

        base_summary_row = last_emp_row + 1

        for s_idx, (cat, label, criteria, target_rows) in enumerate(summary_labels):
            s_row = base_summary_row + s_idx

            # Label in Column F
            ws.cell(s_row, 6, cat).alignment = ALIGN_CENTER
            ws.cell(s_row, 6).font = FONT_BOLD
            ws.cell(s_row, 6).fill = SUMMARY_LABEL_FILL
            ws.cell(s_row, 6).border = THIN_BORDER

            # Merge columns A to E for aesthetic cleanliness
            ws.merge_cells(f"A{s_row}:E{s_row}")
            for c in range(1, 6):
                ws.cell(s_row, c).border = THIN_BORDER

            # Daily formulas for columns G through last_day_col
            for day_idx in range(num_days):
                col_idx = 7 + day_idx
                c_letter = get_column_letter(col_idx)

                if criteria == "SUM":
                    # Sum Standby + Normal + Off
                    sum_start = base_summary_row
                    sum_end = base_summary_row + 2
                    formula = f"=SUM({c_letter}{sum_start}:{c_letter}{sum_end})"
                else:
                    if target_rows and len(target_rows) > 0:
                        first_r = target_rows[0]
                        last_r = target_rows[-1]
                        formula = f'=COUNTIF({c_letter}{first_r}:{c_letter}{last_r}, "{criteria}")'
                    else:
                        formula = f'=COUNTIF({c_letter}{start_emp_row}:{c_letter}{last_emp_row}, "{criteria}")'

                cell = ws.cell(s_row, col_idx, formula)
                cell.alignment = ALIGN_CENTER
                cell.font = FONT_BOLD
                cell.border = THIN_BORDER
                cell.fill = SUMMARY_LABEL_FILL

        # 5. Column Widths
        column_widths = {
            'A': 6,
            'B': 24,
            'C': 12,
            'D': 18,
            'E': 20,
            'F': 14
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        # Day columns width
        for c in range(7, last_day_col + 6):
            ws.column_dimensions[get_column_letter(c)].width = 5.5
        # Summary cols width
        for c in range(last_day_col + 1, last_day_col + 6):
            ws.column_dimensions[get_column_letter(c)].width = 12

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def import_excel_roster(file_bytes: bytes) -> Dict[str, Any]:
        """
        Parses an uploaded roster excel file and extracts employee info & shifts.
        """
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active

        # Find header rows and employees
        employees = []
        shifts_found = {}

        # Look for employee rows starting from row 4
        emp_id = 1
        for r in range(4, ws.max_row + 1):
            name_val = ws.cell(r, 2).value
            if not name_val or str(name_val).strip() in ['Standby', 'Normal', 'Off', 'Total']:
                break

            emp = {
                "id": emp_id,
                "no": ws.cell(r, 1).value or emp_id,
                "name": str(name_val).strip(),
                "sn": str(ws.cell(r, 3).value or "").strip(),
                "position": str(ws.cell(r, 4).value or "").strip(),
                "location": str(ws.cell(r, 5).value or "").strip(),
                "phone": str(ws.cell(r, 6).value or "").strip(),
                "pattern_id": "night_rotation_15" if "malam" in str(ws.cell(r, 2).value).lower() else "regular_5_2",
                "ref_date": "2026-10-01",
                "ref_cycle_index": 0
            }
            employees.append(emp)

            # Read shifts
            for c in range(7, min(38, ws.max_column + 1)):
                val = ws.cell(r, c).value
                if val is not None:
                    day_num = c - 6
                    # Key as day
                    shifts_found[f"{emp_id}:{day_num}"] = str(val).strip().upper()

            emp_id += 1

        return {
            "employees": employees,
            "shifts": shifts_found
        }
